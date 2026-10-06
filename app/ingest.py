"""Leitura de PDF, DOCX e planilhas -> Documents do LangChain com metadados de origem."""
import io
from pathlib import Path

import pandas as pd
from docx import Document as DocxFile
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)
SUPPORTED = {".pdf", ".docx", ".xlsx", ".xls", ".csv"}


def _pdf(data: bytes, name: str):
    for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            yield Document(page_content=text, metadata={"source": name, "location": f"p. {i}"})


def _docx(data: bytes, name: str):
    doc = DocxFile(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    for t in doc.tables:
        for row in t.rows:
            text += "\n" + " | ".join(c.text.strip() for c in row.cells)
    if text.strip():
        yield Document(page_content=text, metadata={"source": name, "location": "documento"})


def _sheet(df: pd.DataFrame, name: str, sheet: str):
    df = df.dropna(how="all").fillna("")
    lines = [" | ".join(f"{c}: {v}" for c, v in row.items() if v != "") for _, row in df.iterrows()]
    if lines:
        yield Document(page_content="\n".join(lines), metadata={"source": name, "location": f"aba {sheet}"})


def _sheets(data: bytes, name: str, ext: str):
    if ext == ".csv":
        yield from _sheet(pd.read_csv(io.BytesIO(data)), name, "csv")
    else:
        for sheet, df in pd.read_excel(io.BytesIO(data), sheet_name=None).items():
            yield from _sheet(df, name, sheet)


def load_chunks(name: str, data: bytes) -> list[Document]:
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Formato não suportado: {ext}")
    docs = {".pdf": lambda: _pdf(data, name), ".docx": lambda: _docx(data, name)}.get(
        ext, lambda: _sheets(data, name, ext)
    )()
    return splitter.split_documents(list(docs))
